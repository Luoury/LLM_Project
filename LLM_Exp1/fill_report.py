# -*- coding: utf-8 -*-
"""
按实验报告样板填写最终报告：
- 填充环境表、基线配置表、结果汇总表、AI 协作记录表、复测表
- 在图片占位单元格插入 result_*.png 曲线图
- 填写各节分析文字与思考题作答（数据全部来自 experiment_results.json 真实运行）
"""
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE = "/home/luosury/.trae-cn-server/data/User/workspaceStorage/7ae190af278cceef3cb5d7cf616b8667/external-files/6aa0bedb6509f4fc65a8718c/0-mttpr5tq-ewfc/作业一_实验报告样板.docx"
OUT = "学号_姓名_实验作业一.docx"

data = json.load(open("experiment_results.json", encoding="utf-8"))
R = {r["key"]: r for r in data["results"]}
fail = data["fail"]
scan = data["lr_scan"]

doc = Document(TEMPLATE)


# ---------- 工具函数 ----------
def set_cell(cell, text, bold=False):
    """覆写单元格文字（保留单元格属性）。"""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.bold = bold


def insert_image_cell(tbl, img_path, width=6.0):
    cell = tbl.cell(0, 0)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img_path, width=Inches(width))


def is_blank(p):
    t = p.text.strip()
    return bool(t) and set(t) <= set("_＿ ")


def replace_para(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.size = Pt(10.5)


# ---------- 1. 封面日期 ----------
for p in doc.paragraphs:
    if p.text.strip().startswith("完成日期"):
        replace_para(p, "完成日期：2026 年 9 月 9 日")
        break

# ---------- 2. 删除“提示：……”段落 ----------
for p in list(doc.paragraphs):
    if p.text.strip().startswith("提示："):
        p._element.getparent().remove(p._element)

# ---------- 3. 正文空白段落分组并填文字 ----------
BLANK_GROUPS = [
    # 1.1 实验目的
    ["掌握人工神经元与多层感知器（MLP）的基本结构，理解加权和、偏置与激活函数在模型中的作用；掌握神经网络“前向传播—损失计算—反向传播—参数更新”的完整训练三步循环。",
     "以 sklearn 内置手写数字数据集 load_digits（8×8 灰度图、1797 个样本、10 个类别）为对象，采用控制变量法，系统对比激活函数、网络深度/宽度、优化器、学习率、L2 正则、Dropout 与 BatchNorm 对模型性能的影响。",
     "学会使用 TRAE IDE 的 Builder / Chat / 内联补全辅助代码编写、调试与实验分析，建立“AI 生成代码为起点、人来理解决策”的协作范式；养成固定随机种子、如实记录每一组实验数据的习惯。"],
    # 2.3 基线现象观察
    [f"基线训练损失从 {R['baseline']['loss_curve'][0]:.2f} 极其缓慢地下降到 {R['baseline']['loss_curve'][-1]:.2f}（随机猜测水平为 ln10≈2.30）；测试准确率在前 19 个 epoch 一直停在 10.0%，30 个 epoch 结束仅为 22.22%。",
     "网络确实在学习（损失单调下降、准确率后期缓慢爬升），但远未达到可用水平，符合指南“能学但学不好”的预期。",
     "原因有三：① Sigmoid 导数最大仅 0.25，梯度反向传播时层层衰减（梯度消失）；② SGD 固定学习率 0.1，且全批量训练每个 epoch 只更新 1 次参数，30 个 epoch 仅 30 次更新，步数太少；③ 单隐藏层 32 个神经元的容量也有限。"],
    # 3.1 实验1 ReLU
    ["动机：Sigmoid 导数最大只有 0.25，多层连乘导致梯度消失；ReLU 在正输入区间导数恒为 1，梯度可无损回传，通常能显著加快收敛。本组仅将激活函数由 Sigmoid 改为 ReLU，其余配置与基线完全一致。",
     "结果：测试准确率从 22.22% 提升到 47.50%（+25.3 个百分点），准确率曲线约从第 15 个 epoch 开始明显抬升，收敛速度显著快于基线；但受 SGD 与极少更新步数的限制，最终仍停留在中等水平。"],
    # 3.2 实验2 加深
    ["动机：增大模型容量（1 层×32 → 2 层×128）预期增强拟合能力。本组严格保持 Sigmoid + SGD(0.1) 不变，只隔离“加深网络”这一个变量。",
     "结果：准确率不升反降到 7.22%，比随机猜测（10%）还差。原因是网络加深后 Sigmoid 梯度消失随层数指数级加剧（两层连乘梯度上限约 0.25²≈0.06），第一层几乎接收不到有效梯度；同时小数据集上参数量远超样本约束，初始化时大量神经元进入饱和区。这说明“更大更深”必须配合合适的激活函数与优化手段才有效。"],
    # 3.3 实验3 Adam
    ["动机：SGD 对所有参数使用同一固定学习率且无动量，在损失曲面平缓区域更新缓慢；Adam 为每个参数维护梯度一阶矩（动量）与二阶矩（梯度平方）的滑动平均，自适应地调整每个参数的有效学习率。本组仅把优化器换成 Adam（lr=0.01）。",
     "结果：测试准确率跃升至 82.22%（+60.0 个百分点），是单项改进中提升最大的之一，损失从 2.30 降到 1.24。说明在本实验“更新步数少”的场景下，优化器选择对结果影响极大。"],
    # 3.4 实验4 lr=0.01
    ["动机：按指南把学习率从 0.1 调为 0.01，观察学习率过小的现象（优化器仍为 SGD，其余不变）。",
     "结果：测试准确率仅 10.00%，比基线还差；训练 30 个 epoch 后损失仍为 2.33，几乎等于初始值。步长缩小 10 倍后，本就只有 30 次的参数更新几乎没有离开初始点。说明学习率好坏必须结合优化器与更新步数评价：本实验全批量设置下 0.1 对 SGD 已属偏小，0.01 直接导致欠拟合。"],
    # 3.5 实验5 L2
    ["动机：L2 正则（weight_decay）通过惩罚大权重抑制过拟合、缩小训练与测试的差距。本组在基线配置上加入 wd=1e-4。",
     "结果：测试准确率为 22.22%，与基线持平。原因是基线本身处于欠拟合（高偏差）状态——模型连训练集都远未拟合好，此时惩罚权重反而进一步压缩了模型表达能力。正则化解决的是高方差（过拟合）问题，用在欠拟合阶段没有收益。"],
    # 3.6 实验6 Dropout
    ["动机：Dropout 在训练时以概率 p=0.2 随机失活神经元，减少神经元间的共适应，预期增强泛化能力。",
     "结果：测试准确率为 19.72%，比基线略低，训练损失曲线也略高于基线。与 L2 同理：基线欠拟合时模型容量本就不足，Dropout 等效于进一步缩小有效网络规模并引入训练噪声，拖慢了本就缓慢的收敛。Dropout 应在容量充足、出现过拟合迹象时使用。"],
    # 3.7 实验7 BN
    ["动机：BatchNorm 在每个隐藏层后对该批激活做标准化（再用可学习参数缩放平移），稳定各层输入分布，缓解内部协变量偏移，减少 Sigmoid 饱和。本组仅插入 BatchNorm1d，其余不变。",
     "结果：测试准确率从 22.22% 大幅提升到 79.17%（+57.0 个百分点），与换 Adam 的提升相当；损失稳定下降到 1.93。BN 使隐层输入保持在 Sigmoid 的非饱和敏感区，梯度信号明显增强——这解释了为何 BN 能让“差激活函数 + 差优化器”的组合也能训练起来。"],
    # 3.8 实验8 最优组合
    ["最优配置：ReLU + 2 个隐藏层×128 神经元 + BatchNorm + Adam(lr=0.001)，30 epoch 全批量，其余与基线一致。值得注意的是：加深网络在 Sigmoid 下是灾难（实验2，7.22%），但在 ReLU+BN 配合下容量优势得以发挥。",
     "结果：固定随机种子复测 3 次（seed=42/43/44），测试准确率分别为 97.22%、94.44%、95.56%，平均 95.74%，显著优于所有单项改进（最高单项为 Adam 的 82.22%），相对基线提升约 73.5 个百分点，平均训练耗时仅 0.17s。"],
    # 5.2 失败实验分析
    [f"现象：lr=10.0 时第 2 个 epoch 损失从 {fail['loss_curve'][0]:.2f} 跳到 {fail['loss_curve'][1]:.2f}（不降反升），此后损失在 2.0~2.8 之间反复震荡（如 2.11→2.59→2.29、2.06→2.56），测试准确率在 10%~37% 之间剧烈跳动，无法稳定收敛，最终停在 36.94%。",
     "原理：参数更新量为 w←w−η·∂L/∂w，学习率过大会使单步更新越过损失最小值，在损失曲面“峡谷”两侧来回弹射。本组未出现 NaN/Inf，是因为 Sigmoid 饱和后梯度趋近于 0，对步长有一定“自限”作用；但参数一旦被甩进饱和区便难以退出，模型始终学不好。",
     f"学习率扫描（加分项）：在基线配置上扫描 lr=0.1/1.0/5.0/10.0/20.0，最终准确率分别为 {scan[0]['acc']:.2f}%/{scan[1]['acc']:.2f}%/{scan[2]['acc']:.2f}%/{scan[3]['acc']:.2f}%/{scan[4]['acc']:.2f}%。lr=1.0 反而远好于 0.1——全批量训练 30 个 epoch 只有 30 次参数更新，小学习率根本走不动；发散临界点出现在 5.0~10.0 之间，lr=20 时训练基本失败。"],
    # 5.3 总结论
    ["本实验设置下各设计维度的有效性排序为：优化器（Adam）与 BatchNorm 提升最大（约 +57~60 个百分点），激活函数（ReLU）次之（+25 个百分点），学习率必须与优化器和更新步数匹配；L2 正则与 Dropout 在欠拟合的基线上没有收益。",
     "改进不是简单叠加而是有条件的组合：加深网络在 Sigmoid 下使准确率暴跌到 7.22%，但同样的 2×128 结构配合 ReLU+BN+Adam 后达到 95% 以上——容量优势需要良好的优化条件才能释放。",
     "小数据集上存在偏差-方差权衡：基线欠拟合（高偏差）时正则化反而有害；组合模型容量大幅提升后，BN 与较短训练起到了稳定与正则作用，测试准确率稳定在 94%~97%。",
     "方法体会：控制变量法（每组只改一处）、共用数据划分与随机种子、如实记录“变差”的实验（实验2、实验4、失败组都是有价值的结论），是得到可信实验结论的前提。"],
    # 思考题 1
    ["现象：实验1 仅把 Sigmoid 换成 ReLU，测试准确率即从 22.22% 升到 47.50%，且曲线约从第 15 个 epoch 起快速抬升，收敛明显加快。",
     "数学解释：Sigmoid σ(x)=1/(1+e⁻ˣ) 的导数为 σ(x)(1−σ(x))，最大值仅 0.25（x=0 处），两端趋近于 0。反向传播按链式法则连乘，经过 k 层后梯度幅度至多衰减为 0.25ᵏ，浅层参数几乎得不到更新信号，即梯度消失。",
     "ReLU 在 x>0 时导数恒为 1，正区间梯度无损传递，深层网络也能获得有效梯度，且计算简单、激活稀疏，因此收敛更快。其副作用是 x<0 时梯度为 0（“神经元死亡”风险），本实验数据简单未暴露该问题。"],
    # 思考题 2
    ["影响不相同。SGD 对所有参数使用统一的全局学习率：把 lr 从 0.1 降到 0.01，每个参数的更新步长都直接缩小 10 倍；本实验只有 30 次更新，步长缩小后参数几乎停在初始点（实验4 准确率仅 10%）。",
     "Adam 的实际更新量约为 η·m̂/(√v̂+ε)：m̂ 是梯度一阶矩（动量），v̂ 是梯度平方的二阶矩。分母 √v̂ 对每个参数做自适应缩放——梯度大的参数有效步长被压小，梯度长期小的参数相对步长更大。",
     "因此 Adam 对全局学习率的选择远不敏感（常用 1e-3 即可适配多数任务），lr 改变 10 倍主要影响整体节奏而非“走不走得动”；SGD 没有这种自调节机制，学习率是决定性的。"],
    # 思考题 3
    ["内部协变量偏移（Internal Covariate Shift）指：训练中每层参数更新会改变下一层输入的分布，分布持续漂移迫使各层不断重新适应；大学习率会加剧这种漂移，使激活容易进入 Sigmoid/tanh 的饱和区，导致梯度趋零、训练震荡。",
     "BatchNorm 在每个小批量数据上对隐层预激活做标准化（均值 0、方差 1），再通过可学习的 γ、β 缩放平移，把各层输入分布“锚定”在稳定范围内，显著削弱了参数更新对下游分布的冲击。",
     "分布稳定后损失曲面更平滑、梯度更大且方向更一致，大学习率造成的更新不会把网络推入饱和区，因而 BN 能让模型容忍更大学习率、收敛更快更稳（实验7 中 BN 使 Sigmoid+SGD 组合达到 79.17%）。"],
    # 思考题 4
    ["PyTorch 的 nn.Dropout 采用“反向 Dropout”（inverted dropout）实现：训练时以概率 p 随机将神经元输出置 0，同时对保留下来的输出乘以 1/(1−p) 进行放大。",
     "这样训练期神经元输出的期望幅度与无 Dropout 时一致：E[output]=(1−p)·(x/(1−p))+p·0 = x。",
     "推理时模型调用 model.eval() 切换到评估模式，Dropout 层变为恒等映射（不丢弃、也不缩放）；由于训练时的 1/(1−p) 放大已经补偿了被丢弃神经元的期望贡献，推理输出的尺度自然与训练保持一致，无需任何额外处理。"],
    # 思考题 5
    ["偏差-方差权衡：模型容量不足时偏差主导（欠拟合），训练误差与测试误差都高；容量过大时方差主导（过拟合）——参数量远超训练样本能提供的约束，网络会记忆训练样本中的噪声，训练误差很低但测试误差反弹。",
     "本数据集仅 1437 个训练样本，而 2×128 网络约有 2.5 万个参数，样本约束严重不足；实验2 中它还叠加了深层 Sigmoid 梯度消失的优化困难，准确率仅 7.22%。",
     "因此“更大更深”不是免费的：容量优势需要 ReLU 等良好激活、BatchNorm 稳定化、Adam 等高效优化器以及必要时的正则化配合，才能转化为泛化收益；否则只是增加过拟合风险与优化难度。"],
    # 思考题 6
    ["本次实验中 TRAE 最有价值的环节：用 Builder 一次性生成符合指南要求的实验框架（可配置 MLP 类、训练循环、绘图），用 Chat 快速解释报错与概念（如 RTX 50 系 Blackwell 架构需安装 cu128 版 PyTorch、深层 Sigmoid 为何退化），用内联补全快速生成训练三步与评估代码，显著减少了重复劳动。",
     "但关键决策必须由人完成：实验组的控制变量设计、判断“准确率下降”是 bug 还是有价值的结论（实验2、实验4）、确定最优组合配置、核对表格中每一个数字都来自真实运行并对其作出解释。",
     "AI 生成的代码是“起点”而非“终点”：读懂、验证、修改之后才成为自己的实验。人负责提出问题、设计验证、承担结论；AI 负责加速执行与提供线索——这种分工正是“人机协作做实验”的核心。"],
]

# 按文档顺序收集连续空白段落分组
groups, cur = [], []
for p in doc.paragraphs:
    if is_blank(p):
        cur.append(p)
    else:
        if cur:
            groups.append(cur)
            cur = []
if cur:
    groups.append(cur)

assert len(groups) == len(BLANK_GROUPS), f"空白组数 {len(groups)} 与内容组数 {len(BLANK_GROUPS)} 不一致"

for blanks, contents in zip(groups, BLANK_GROUPS):
    for i, p in enumerate(blanks):
        if i < len(contents):
            replace_para(p, contents[i])
        else:
            replace_para(p, "")

# ---------- 4. 2.2 基线结果数字 ----------
for p in doc.paragraphs:
    if "最终测试准确率" in p.text and "训练耗时" in p.text:
        replace_para(p, f"最终测试准确率：{R['baseline']['acc']:.2f} %　　　训练耗时：{R['baseline']['time']:.2f} s")
        break

# ---------- 5. 表格填充 ----------
t_env, t_base, t_res, t_ai, t_re = doc.tables[0], doc.tables[1], doc.tables[3], doc.tables[12], doc.tables[13]

# 表1 环境表
env_rows = [
    ("Ubuntu 22.04.5 LTS（x86_64）", ""),
    ("Python 3.10.12", "独立虚拟环境 LLM_Venv，多实验共享"),
    ("PyTorch 2.11.0+cu128", "CUDA 可用：是（NVIDIA GeForce RTX 5060 Laptop GPU）"),
    ("scikit-learn 1.7.2 / NumPy 2.2.6 / Matplotlib 3.10.9", ""),
    ("TRAE IDE（国内版，最新）", "使用的核心功能：Builder 生成框架、Chat 答疑排错、内联补全"),
]
for i, (name, note) in enumerate(env_rows, start=1):
    set_cell(t_env.rows[i].cells[1], name)
    set_cell(t_env.rows[i].cells[2], note)

# 表2 基线配置表
set_cell(t_base.rows[1].cells[1], "64 → 隐藏层 1×32（Sigmoid 激活） → 输出 10")
set_cell(t_base.rows[2].cells[1], "SGD / lr = 0.1")

# 表3 结果汇总表（行序：基线、exp1~8、失败组）
order = ["baseline", "exp1", "exp2", "exp3", "exp4", "exp5", "exp6", "exp7", "exp8"]
conclusions = {
    "baseline": "损失几乎不降，准确率仅 22.22%，“能学但学不好”",
    "exp1": "ReLU 缓解梯度消失，升至 47.50%（+25.3pp）",
    "exp2": "深层 Sigmoid 梯度消失加剧，暴跌至 7.22%（劣于随机猜测）",
    "exp3": "Adam 自适应学习率，跃升至 82.22%（+60.0pp）",
    "exp4": "学习率过小、30 步未收敛，降至 10.00%",
    "exp5": "欠拟合阶段 L2 无收益，22.22% 与基线持平",
    "exp6": "欠拟合阶段 Dropout 无收益，19.72% 略降",
    "exp7": "BN 稳定层输入分布，大幅升至 79.17%（+57.0pp）",
    "exp8": "ReLU+2×128+BN+Adam(0.001)，3 次复测均值 95.74%",
}
for i, key in enumerate(order, start=1):
    r = R[key]
    set_cell(t_res.rows[i].cells[3], f"{r['acc']:.2f}")
    set_cell(t_res.rows[i].cells[4], f"{r['time']:.2f}")
    set_cell(t_res.rows[i].cells[5], conclusions[key])
# 失败组（最后一行）
fail_row = t_res.rows[10]
set_cell(fail_row.cells[2], "学习率过大（lr = 10.0）")
set_cell(fail_row.cells[3], f"{fail['acc']:.2f}")
set_cell(fail_row.cells[4], f"{fail['time']:.2f}")
set_cell(fail_row.cells[5], "损失在 2.0~2.8 间震荡、准确率 10%~37% 跳动，不收敛")

# 表4 AI 协作记录
ai_records = [
    ("Builder 生成实验框架",
     "“请在当前目录创建 mlp_digits.py，用 PyTorch 实现手写数字分类实验框架：load_digits 按 8:2 分层划分（seed=42），标签转 torch.long；MLP 的隐藏层/激活/Dropout/BatchNorm 可配置；优化器、lr、weight_decay 可配置；每 epoch 记录损失与准确率；绘制曲线保存 result_baseline.png。”",
     "AI 生成含数据加载、可配置 MLP、训练循环与绘图的完整框架。我通读核查：标签 .long()、CrossEntropyLoss、分层划分均正确；自行补充 Linux 中文字体回退（Noto Sans CJK SC）、GPU 设备选择，并将基线改回指南规定的全批量 30 epoch。"),
    ("Chat 环境排错",
     "“RTX 5060 Laptop 上 PyTorch 已安装但 CUDA 不可用/该装哪个版本？”",
     "AI 指出 RTX 50 系为 Blackwell 架构（sm_120），需 CUDA 12.8+ 编译的 PyTorch 2.7+。我据此执行 pip install torch --index-url https://download.pytorch.org/whl/cu128，验证 cuda.is_available() 为 True。"),
    ("Chat 结果分析",
     "“把网络加深到 2 层×128（仍用 Sigmoid）后准确率为什么比 1 层 32 还差？”",
     "AI 解释：深层 Sigmoid 梯度连乘消失、小数据集上容量过大、初始化易进入饱和区。我对照实验2 曲线（7.22%）核实后写入报告，并在最优组合中用 ReLU+BN 解决该问题。"),
    ("内联补全 / 智能修复",
     "编写 run() 训练函数与评估代码时触发内联补全；venv 移动后报错时参考智能修复建议。",
     "AI 补全了 zero_grad/backward/step 训练三步与 torch.no_grad() 评估块；我核对后补上 model.train()/eval() 模式切换，并自行加入 torch.manual_seed、use_deterministic_algorithms 与“先设种子再构造模型”，保证结果可复现。"),
]
for i, (scene, prompt, adopt) in enumerate(ai_records, start=1):
    row = t_ai.rows[i]
    set_cell(row.cells[1], scene)
    set_cell(row.cells[2], prompt)
    set_cell(row.cells[3], adopt)

# 表5 最优组合复测
runs = R["exp8"]["runs"]
for i, run in enumerate(runs, start=1):
    row = t_re.rows[i]
    set_cell(row.cells[1], str(run["seed"]))
    set_cell(row.cells[2], f"{run['acc']:.2f}")
    set_cell(row.cells[3], f"{run['time']:.2f}")
avg_row = t_re.rows[4]
set_cell(avg_row.cells[1], "42/43/44")
set_cell(avg_row.cells[2], f"{R['exp8']['acc']:.2f}")
set_cell(avg_row.cells[3], f"{R['exp8']['time']:.2f}")

# ---------- 6. 插入曲线图 ----------
img_map = {
    2: "result_baseline.png",
    4: "result_exp1.png",
    5: "result_exp2.png",
    6: "result_exp3.png",
    7: "result_exp4.png",
    8: "result_exp5.png",
    9: "result_exp6.png",
    10: "result_exp7.png",
    11: "result_exp8_run1.png",
    14: "result_lr_too_big.png",
}
for idx, img in img_map.items():
    insert_image_cell(doc.tables[idx], img)

# ---------- 7. 5.2 节附加：学习率扫描图（图11） ----------
anchor = None
for p in doc.paragraphs:
    if p.text.strip().startswith("图 10"):
        anchor = p
        break
if anchor is not None:
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("图 11  学习率扫描（0.1/1.0/5.0/10.0/20.0）：发散临界点观察（result_lr_scan.png）")
    r.font.bold = True
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture("result_lr_scan.png", width=Inches(5.2))
    # 移动到图10标题之后（先图后题注）
    anchor._p.addnext(cap._p)
    anchor._p.addnext(img_p._p)

doc.save(OUT)
print("报告已生成:", OUT)
